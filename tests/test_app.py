"""
PRUEBAS AUTOMÁTICAS · tests/test_app.py

Estas pruebas simulan un navegador con app.test_client(): hacen las mismas
peticiones que una persona, pero sin abrir la app ni usar el ratón.
"""
import html as html_lib
import io
import os
import re
import sys
from pathlib import Path

import openpyxl
import pytest

# ---------------------------------------------------------------------------
# Credenciales de prueba
# ---------------------------------------------------------------------------
# Se definen ANTES de importar app para que las pruebas no dependan del .env
# de la máquina. load_dotenv() no sobreescribe variables que ya existan.
os.environ["ADMIN_CORREO"] = "admin@prueba.cr"
os.environ["ADMIN_CLAVE"] = "clave-admin-de-prueba"
os.environ["USUARIO_CORREO"] = "usuario@prueba.cr"
os.environ["USUARIO_CLAVE"] = "clave-usuario-de-prueba"
os.environ["SECRET_KEY"] = "llave-secreta-solo-para-pruebas"
# Base de datos vacía y modo Vercel: las pruebas leen los Excel de insumos.
os.environ["DATABASE_URL"] = ""
os.environ["VERCEL"] = "1"

# Permite encontrar app.py aunque pytest se ejecute desde otra carpeta.
CARPETA_RAIZ = Path(__file__).resolve().parent.parent
if str(CARPETA_RAIZ) not in sys.path:
    sys.path.insert(0, str(CARPETA_RAIZ))

from app import app  # noqa: E402  (se importa después de las variables de entorno)


@pytest.fixture()
def cliente():
    """Devuelve un navegador de mentira con su propia sesión y sus propias cookies."""
    app.config.update(TESTING=True)
    return app.test_client()


def _entrar(cliente, correo, clave):
    """Simula escribir el correo y la contraseña y pulsar "Entrar"."""
    return cliente.post("/login", data={"correo": correo, "clave": clave})


def _numero_de_tarjeta(html, etiqueta):
    """Lee el número que muestra una de las tarjetas del tablero."""
    patron = rf'<div class="num">(\d+)</div>\s*<div class="lbl">{re.escape(etiqueta)}</div>'
    coincidencia = re.search(patron, html)
    assert coincidencia, f"No encontré la tarjeta: {etiqueta}"
    return int(coincidencia.group(1))


def test_sin_sesion_redirige_a_login(cliente):
    """Simula a alguien que escribe la dirección de la app sin haber entrado."""
    respuesta = cliente.get("/")

    # 302 significa "vaya a otra página"; el encabezado Location dice cuál.
    assert respuesta.status_code == 302
    assert respuesta.headers["Location"].endswith("/login")


def test_clave_incorrecta_no_entra(cliente):
    """Simula a alguien que escribe bien el correo pero se equivoca en la clave."""
    respuesta = cliente.post(
        "/login",
        data={"correo": os.environ["ADMIN_CORREO"], "clave": "esta-clave-no-es"},
    )

    # Se queda en la pantalla de entrada (200) con el mensaje rojo de error.
    assert respuesta.status_code == 200
    assert "Correo o contraseña incorrectos." in respuesta.get_data(as_text=True)

    # Y sigue sin sesión: al pedir el tablero la app lo devuelve al login.
    assert cliente.get("/").status_code == 302


def test_usuario_no_puede_cargar(cliente):
    """Simula a un usuario con rol "usuario" escribiendo a mano la dirección /cargar."""
    _entrar(cliente, os.environ["USUARIO_CORREO"], os.environ["USUARIO_CLAVE"])

    respuesta = cliente.get("/cargar")

    # 403 significa "entró, pero no tiene permiso para esta sección".
    assert respuesta.status_code == 403
    assert "solo para administradores" in respuesta.get_data(as_text=True).lower()


def test_excel_modificado_cambia_contadores(cliente):
    """Simula al admin subiendo un Registro de Evaluaciones con las notas de
    Camila Rojas Méndez (101230456) cambiadas a 40, sin tocar el archivo original."""
    _entrar(cliente, os.environ["ADMIN_CORREO"], os.environ["ADMIN_CLAVE"])

    ruta_maestro = CARPETA_RAIZ / "insumos" / "Maestro_Estudiantes.xlsx"
    ruta_evaluaciones = CARPETA_RAIZ / "insumos" / "Registro_Evaluaciones.xlsx"

    # Se abre el Excel original en memoria y se modifica una COPIA.
    libro = openpyxl.load_workbook(ruta_evaluaciones)
    hoja = libro.active
    encabezados = [celda.value for celda in hoja[1]]
    columna_id = encabezados.index("Identificacion") + 1
    columna_nota = encabezados.index("Nota") + 1

    filas_camila = 0
    for numero_fila in range(2, hoja.max_row + 1):
        identificacion = str(hoja.cell(numero_fila, columna_id).value).strip()
        if identificacion == "101230456":
            hoja.cell(numero_fila, columna_nota).value = 40
            filas_camila += 1

    assert filas_camila > 0, "La prueba esperaba encontrar evaluaciones de Camila"

    memoria_evaluaciones = io.BytesIO()
    libro.save(memoria_evaluaciones)  # todo queda en memoria, no en el disco
    memoria_evaluaciones.seek(0)

    memoria_maestro = io.BytesIO(ruta_maestro.read_bytes())

    respuesta = cliente.post(
        "/cargar",
        data={
            "accion": "ver",
            "maestro": (memoria_maestro, "Maestro_Estudiantes.xlsx"),
            "evaluaciones": (memoria_evaluaciones, "Registro_Evaluaciones.xlsx"),
        },
        content_type="multipart/form-data",
    )

    assert respuesta.status_code == 200
    html = respuesta.get_data(as_text=True)

    # Con las notas de Camila en 40 pasa de Aprobación a Participación:
    # 19 - 1 = 18 aprobación, 3 + 1 = 4 participación y sigue 1 sin certificado.
    assert _numero_de_tarjeta(html, "Certificados de Aprobación") == 18
    assert _numero_de_tarjeta(html, "Certificados de Participación") == 4
    assert _numero_de_tarjeta(html, "Sin certificado") == 1


def test_certificado_se_descarga(cliente):
    """Simula abrir el primer enlace "Ver PDF" del tablero y comprueba que
    el navegador recibe un PDF de verdad, no una página de error."""
    _entrar(cliente, os.environ["ADMIN_CORREO"], os.environ["ADMIN_CLAVE"])

    tablero = cliente.get("/")
    html = tablero.get_data(as_text=True)

    # Los enlaces ahora usan /certificado?identificacion=...&programa=...
    enlaces = [
        html_lib.unescape(enlace)
        for enlace in re.findall(r'href="(/certificado[^"]*)"', html)
    ]
    assert len(enlaces) == 22

    respuesta = cliente.get(enlaces[0])

    assert respuesta.status_code == 200
    assert respuesta.headers["Content-Type"].startswith("application/pdf")
    assert respuesta.data.startswith(b"%PDF")
